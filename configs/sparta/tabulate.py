import os

import pandas as pd
from openpyxl import Workbook

STATS_ROOT = "stats"

ARCHS = ["Sequential", "Parallel", "Gamma", "SparTA"]

metrics = [
    "simTicks",
    "Speedup achieved by SparTA",
    "Energy (1e-9 J)",
    "EDP (1e-15)",
    "Energy Efficiency",
]

models = sorted(os.listdir(STATS_ROOT))

wb = Workbook()
ws = wb.active


# ------------------------------------------------
# Header row 1
# ------------------------------------------------

ws.cell(row=1, column=1, value="Architecture")

col = 2
for model in models:
    ws.merge_cells(
        start_row=1, start_column=col, end_row=1, end_column=col + 3
    )
    ws.cell(row=1, column=col, value=model)
    col += 4


# ------------------------------------------------
# Header row 2
# ------------------------------------------------

col = 2
for model in models:
    for arch in ARCHS:
        ws.cell(row=2, column=col, value=arch)
        col += 1


# ------------------------------------------------
# Fill metrics rows
# ------------------------------------------------

for i, metric in enumerate(metrics):

    ws.cell(row=3 + i, column=1, value=metric)

    col = 2

    for model in models:

        csv_path = os.path.join(STATS_ROOT, model, "results.csv")

        if not os.path.exists(csv_path):
            for _ in ARCHS:
                ws.cell(row=3 + i, column=col, value="")
                col += 1
            continue

        df = pd.read_csv(csv_path).set_index("Architecture")

        spartaticks = df.loc["SparTA", "simTicks"]
        seq_energy = df.loc["Sequential", "Energy"]

        for arch in ARCHS:

            if arch not in df.index:
                ws.cell(row=3 + i, column=col, value="")
                col += 1
                continue

            ticks = df.loc[arch, "simTicks"]
            energy = df.loc[arch, "Energy"] * 1e9
            edp = df.loc[arch, "EDP"] * 1e15

            speedup = ticks / spartaticks if spartaticks != 0 else ""
            efficiency = seq_energy / (energy * 1e-9) if energy != 0 else ""

            value_map = {
                "simTicks": round(ticks, 2),
                "Speedup achieved by SparTA": round(speedup, 2),
                "Energy (1e-9 J)": round(energy, 2),
                "EDP (1e-15)": round(edp, 2),
                "Energy Efficiency": round(efficiency, 2),
            }

            ws.cell(row=3 + i, column=col, value=value_map[metric])

            col += 1


wb.save("sparta_table.xlsx")

print("Saved: sparta_table.xlsx")
